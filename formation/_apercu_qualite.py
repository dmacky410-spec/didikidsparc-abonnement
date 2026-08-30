# -*- coding: utf-8 -*-
"""Rendu HTML fidèle des diapositives, à partir de la géométrie réelle du .pptx."""
import html
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

EMU = 914400.0
PX = 96.0
prs = Presentation("formation/Didikids_Parc_Guide_Agent.pptx")
SW = prs.slide_width / EMU * PX
SH = prs.slide_height / EMU * PX

ALIGN = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right"}


def color_of(fmt):
    try:
        if fmt.type is not None and fmt.fore_color.type is not None:
            return "#%02X%02X%02X" % tuple(fmt.fore_color.rgb)
    except Exception:
        pass
    return None


SCALE = 0.285
out = ["<style>body{background:#2b2b2b;font-family:Calibri,Carlito,sans-serif;margin:0;"
       "padding:14px;display:flex;flex-wrap:wrap;gap:16px;justify-content:center}"
       ".cell{width:%dpx}"
       ".lbl{color:#fff;font:700 13px sans-serif;margin:0 0 4px}"
       ".wrap{width:%dpx;height:%dpx;overflow:hidden;box-shadow:0 3px 12px rgba(0,0,0,.6)}"
       ".slide{position:relative;width:%dpx;height:%dpx;background:#fff;overflow:hidden;"
       "transform:scale(%s);transform-origin:top left}"
       ".sh{position:absolute;box-sizing:border-box}</style>"
       % (SW*SCALE, SW*SCALE, SH*SCALE, SW, SH, SCALE)]

for i, slide in enumerate(prs.slides, 1):
    out.append(f'<div class="cell"><div class="lbl">Diapositive {i}</div>'
               f'<div class="wrap"><div class="slide">')
    for sh in slide.shapes:
        x = sh.left / EMU * PX
        y = sh.top / EMU * PX
        w = sh.width / EMU * PX
        h = sh.height / EMU * PX
        style = f"left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;height:{h:.1f}px;"
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
            fill = color_of(sh.fill)
            if fill:
                style += f"background:{fill};"
            try:
                if sh.line.fill.type is not None and sh.line.color and sh.line.color.rgb:
                    style += "border:3px solid #%02X%02X%02X;" % tuple(sh.line.color.rgb)
            except Exception:
                pass
            name = str(sh.shape_type)
            try:
                if "OVAL" in str(sh.auto_shape_type):
                    style += "border-radius:50%;"
                elif "ROUNDED" in str(sh.auto_shape_type):
                    style += "border-radius:14px;"
            except Exception:
                pass
        inner = ""
        if sh.has_text_frame and sh.text_frame.text.strip():
            tf = sh.text_frame
            va = str(tf.vertical_anchor or "")
            paras = []
            for p in tf.paragraphs:
                runs = []
                for r in p.runs:
                    fs = r.font.size.pt if r.font.size else 16
                    col = "#000"
                    try:
                        if r.font.color and r.font.color.rgb:
                            col = "#%02X%02X%02X" % tuple(r.font.color.rgb)
                    except Exception:
                        pass
                    bold = "700" if r.font.bold else "400"
                    runs.append(f'<span style="font-size:{fs}px;color:{col};'
                                f'font-weight:{bold}">{html.escape(r.text)}</span>')
                al = ALIGN.get(p.alignment, "left")
                sa = p.space_after.pt if p.space_after else 0
                ls = p.line_spacing if isinstance(p.line_spacing, float) else 1.15
                paras.append(f'<div style="text-align:{al};margin-bottom:{sa}px;'
                             f'line-height:{ls}">{"".join(runs) or "&nbsp;"}</div>')
            just = "center" if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE else "flex-start"
            inner = (f'<div style="display:flex;flex-direction:column;height:100%;'
                     f'justify-content:{just}">{"".join(paras)}</div>')
        out.append(f'<div class="sh" style="{style}">{inner}</div>')
    out.append("</div></div></div>")

with open("formation/_preview.html", "w", encoding="utf-8") as f:
    f.write("\n".join(out))
print("rendu ecrit : formation/_preview.html")
