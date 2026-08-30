# -*- coding: utf-8 -*-
"""Guide de l'agent d'accueil — PowerPoint avec reproductions des écrans.

    python3 formation/build_guide_agent.py

Destiné aux employés de base uniquement (pas au gérant).
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE

# ----------------------------------------------------------------- palette
GREEN_DARK = RGBColor(0x2A, 0x7A, 0x3B)
GREEN_MAIN = RGBColor(0x4C, 0xAF, 0x50)
GREEN_PALE = RGBColor(0xE8, 0xF7, 0xEA)
YELLOW = RGBColor(0xF5, 0xC5, 0x18)
YELLOW_PALE = RGBColor(0xFF, 0xF8, 0xE1)
PURPLE = RGBColor(0x7B, 0x2F, 0xBE)
PURPLE_PALE = RGBColor(0xF3, 0xE5, 0xFF)
RED = RGBColor(0xD9, 0x46, 0x3E)
RED_PALE = RGBColor(0xFD, 0xEC, 0xEA)
TEXT = RGBColor(0x2D, 0x3A, 0x2E)
MUTED = RGBColor(0x7C, 0x8F, 0x7E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xD3, 0xE4, 0xD5)
GREY = RGBColor(0xF4, 0xF7, 0xF4)

FONT = "Calibri"
W, H = 13.333, 7.5


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # ------------------------------------------------------------- socle
    def slide(bg=WHITE):
        s = prs.slides.add_slide(blank)
        r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                               prs.slide_width, prs.slide_height)
        r.fill.solid()
        r.fill.fore_color.rgb = bg
        r.line.fill.background()
        r.shadow.inherit = False
        return s

    def text(s, x, y, w, h, runs, size=15, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, space=5, line=None, anchor=MSO_ANCHOR.TOP):
        box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        paras = runs if isinstance(runs, list) else [runs]
        eclates = []
        for p_ in paras:
            eclates.extend(p_.split("\n")) if isinstance(p_, str) and "\n" in p_ \
                else eclates.append(p_)
        for i, para in enumerate(eclates):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(space)
            if line:
                p.line_spacing = line
            for content, opt in (para if isinstance(para, list) else [(para, {})]):
                if content == "":
                    continue
                r = p.add_run()
                r.text = content
                f = r.font
                f.name = FONT
                f.size = Pt(opt.get("size", size))
                f.bold = opt.get("bold", bold)
                f.color.rgb = opt.get("color", color)
        return box

    def box(s, x, y, w, h, fill=WHITE, outline=None, radius=0.08,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE, dashed=False, lw=1.5):
        sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        if outline:
            sh.line.color.rgb = outline
            sh.line.width = Pt(lw)
            if dashed:
                sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            try:
                sh.adjustments[0] = radius
            except (IndexError, KeyError):
                pass
        return sh

    def pastille(s, x, y, d, fill, label, lc=WHITE, size=17):
        c = box(s, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = lc
        return c

    def bouton(s, x, y, w, h, label, fill=GREEN_DARK, lc=WHITE, size=11):
        b = box(s, x, y, w, h, fill=fill, radius=0.5)
        tf = b.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = lc
        return b

    def champ(s, x, y, w, label, valeur="", h=0.38, dashed=False,
              vc=TEXT, lsize=9.5, vsize=11.5):
        text(s, x, y, w, 0.2, label, size=lsize, bold=True, color=MUTED, space=0)
        f = box(s, x, y + 0.24, w, h, fill=WHITE,
                outline=GREEN_MAIN if dashed else BORDER,
                radius=0.25, dashed=dashed, lw=2 if dashed else 1.25)
        if valeur:
            text(s, x + 0.14, y + 0.24 + (h - 0.2) / 2, w - 0.28, 0.22, valeur,
                 size=vsize, color=vc, space=0)
        return y + 0.24 + h

    def entete(s, titre, sous=None):
        text(s, 0.75, 0.5, 11.8, 0.8, titre, size=36, bold=True, color=GREEN_DARK)
        if sous:
            text(s, 0.75, 1.35, 11.8, 0.45, sous, size=16, color=MUTED)

    def etape(s, x, y, n, titre, desc=None, couleur=GREEN_DARK, lc=WHITE, wt=4.2):
        pastille(s, x, y, 0.5, couleur, str(n), lc=lc, size=15)
        text(s, x + 0.72, y + 0.02, wt, 0.35, titre, size=15.5, bold=True, color=TEXT)
        if desc:
            text(s, x + 0.72, y + 0.42, wt, 0.55, desc, size=12.5, color=MUTED, line=1.15)

    def notes(s, t):
        s.notes_slide.notes_text_frame.text = t

    # --------------------------------------------- maquette : fenêtre du logiciel
    def fenetre(s, x, y, w, h, page_active="Accueil / Entrées"):
        box(s, x, y, w, h, fill=WHITE, outline=BORDER, radius=0.04, lw=1.25)
        sw = 1.75
        box(s, x, y, sw, h, fill=GREEN_DARK, radius=0.05)
        text(s, x + 0.16, y + 0.18, sw - 0.3, 0.25, "Didikids Parc",
             size=11.5, bold=True, color=WHITE, space=0)
        menus = ["Accueil / Entrées", "Membres", "Relances WhatsApp",
                 "Historique visites"]
        my = y + 0.62
        for m in menus:
            if m == page_active:
                box(s, x + 0.1, my - 0.06, sw - 0.2, 0.34, fill=YELLOW, radius=0.3)
                text(s, x + 0.2, my + 0.02, sw - 0.4, 0.22, m, size=9.5,
                     bold=True, color=TEXT, space=0)
            else:
                text(s, x + 0.2, my + 0.02, sw - 0.4, 0.22, m, size=9.5,
                     color=GREEN_PALE, space=0)
            my += 0.44
        return x + sw

    # ============================================================== 1. titre
    s = slide(GREEN_DARK)
    text(s, 1.0, 2.15, 11.3, 1.1,
         [[("Didikids ", {"color": WHITE}), ("Parc", {"color": YELLOW})]],
         size=54, bold=True)
    text(s, 1.0, 3.45, 11.3, 0.7, "GUIDE DE L'AGENT D'ACCUEIL",
         size=26, bold=True, color=GREEN_PALE)
    text(s, 1.0, 4.4, 11.3, 0.5,
         "Tout ce que vous ferez au comptoir, écran par écran",
         size=18, color=WHITE)
    box(s, 1.0, 5.5, 5.6, 0.72, fill=YELLOW)
    text(s, 1.3, 5.72, 5.1, 0.35, "Formation interne · Équipe d'accueil",
         size=14, bold=True, color=TEXT)
    notes(s, "Guide destine aux agents d'accueil. Le gerant a sa propre formation.")

    # ============================================================== 2. journée
    s = slide()
    entete(s, "Votre journée en trois gestes",
           "Ce guide couvre exactement ces trois situations")
    items = [
        ("1", "Un enfant arrive", GREEN_DARK, GREEN_PALE,
         "Il pose sa carte, vous lisez la couleur de l'écran. C'est tout."),
        ("2", "Un nouveau client s'inscrit", YELLOW, YELLOW_PALE,
         "Vous créez sa fiche, vous lui donnez une carte, vous encaissez. "
         "Vous faites tout, seul."),
        ("3", "Vous appelez les parents", PURPLE, PURPLE_PALE,
         "Abonnements qui se terminent et anniversaires à venir, "
         "avec le message déjà écrit."),
    ]
    x = 0.75
    for num, titre, coul, fond, desc in items:
        box(s, x, 2.2, 3.75, 4.4, fill=fond)
        pastille(s, x + 0.4, 2.7, 0.85, coul, num,
                 lc=TEXT if coul == YELLOW else WHITE, size=22)
        text(s, x + 0.4, 3.9, 3.0, 0.45, titre, size=20, bold=True, color=GREEN_DARK)
        text(s, x + 0.4, 4.55, 2.95, 1.9, desc, size=14.5, line=1.25)
        x += 4.0
    notes(s, "Annoncer le plan : les entrees, l'inscription, les relances.")

    # ============================================================== 3. connexion
    s = slide()
    entete(s, "Se connecter", "Votre compte est personnel")
    y = 2.3
    for n, t, d in [
        ("1", "Ouvrez le logiciel",
         "Double-cliquez sur l'icône « Didikids Parc » du Bureau. "
         "Une fenêtre noire s'ouvre : ne la fermez jamais."),
        ("2", "Saisissez vos identifiants",
         "Votre nom d'utilisateur et votre mot de passe, remis par le responsable."),
        ("3", "L'écran d'accueil s'ouvre",
         "Vous êtes prêt à recevoir les enfants."),
    ]:
        etape(s, 0.85, y, n, t, d, wt=6.2)
        y += 1.4
    box(s, 8.05, 2.3, 4.5, 3.85, fill=YELLOW_PALE, outline=YELLOW, lw=2.25)
    text(s, 8.45, 2.7, 3.7, 0.4, "RÈGLE IMPORTANTE", size=15, bold=True)
    text(s, 8.45, 3.3, 3.7, 2.5,
         "Ne prêtez votre compte à personne, même à un collègue, même cinq minutes.\n\n"
         "Chaque entrée validée et chaque encaissement portent votre nom.\n\n"
         "C'est votre protection : personne ne pourra vous reprocher l'erreur "
         "d'un autre.",
         size=14, line=1.3)
    notes(s, "Insister sur le compte personnel des le debut.")

    # ============================================================== 4. écran d'accueil
    s = slide()
    entete(s, "L'écran d'accueil", "Voici l'écran devant lequel vous passerez la journée")
    cx = fenetre(s, 0.75, 2.05, 7.6, 4.5)
    text(s, cx + 0.25, 2.28, 3.0, 0.3, "Contrôle des entrées", size=15,
         bold=True, color=GREEN_DARK, space=0)
    box(s, cx + 0.25, 2.75, 5.3, 2.05, fill=WHITE, outline=BORDER, radius=0.1)
    text(s, cx + 0.35, 3.05, 5.1, 0.3, "Présentez la carte", size=15,
         bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER, space=0)
    text(s, cx + 0.35, 3.45, 5.1, 0.3,
         "Le passage est validé automatiquement", size=10.5, color=MUTED,
         align=PP_ALIGN.CENTER, space=0)
    b = box(s, cx + 0.9, 3.9, 4.0, 0.55, fill=GREEN_PALE, outline=GREEN_MAIN,
            dashed=True, radius=0.2, lw=2)
    text(s, cx + 0.9, 4.05, 4.0, 0.3, "ou saisir l'UID ici", size=12,
         bold=True, color=GREEN_DARK, align=PP_ALIGN.CENTER, space=0)
    box(s, cx + 0.25, 5.0, 5.3, 1.3, fill=GREY, radius=0.1)
    text(s, cx + 0.45, 5.2, 4.9, 0.25, "Passages du jour", size=12, bold=True,
         color=GREEN_DARK, space=0)
    ly = 5.6
    for heure, nom, ok in [("14:32", "Aminata Diallo", True),
                           ("14:05", "Sekou Conde", True)]:
        text(s, cx + 0.45, ly, 0.7, 0.22, heure, size=10, color=TEXT, space=0)
        text(s, cx + 1.25, ly, 2.4, 0.22, nom, size=10, bold=True, color=TEXT, space=0)
        box(s, cx + 4.2, ly - 0.03, 0.55, 0.26, fill=GREEN_PALE, radius=0.4)
        text(s, cx + 4.2, ly + 0.02, 0.55, 0.2, "OK", size=8.5, bold=True,
             color=GREEN_DARK, align=PP_ALIGN.CENTER, space=0)
        ly += 0.35
    y = 2.3
    for n, t, d in [
        ("1", "Le menu à gauche", "Vos quatre pages. Si un menu manque, c'est normal : "
         "chacun voit ce qui le concerne."),
        ("2", "La zone de lecture", "Le curseur y reste tout seul. Ne cliquez nulle part, "
         "posez la carte."),
        ("3", "Les passages du jour", "Qui est entré, à quelle heure."),
    ]:
        etape(s, 8.65, y, n, t, d, wt=3.9)
        y += 1.4
    notes(s, "Montrer le vrai ecran en parallele si possible.")

    # ============================================================== 5. valider
    s = slide()
    entete(s, "Valider une entrée", "Le geste que vous répéterez cent fois par jour")
    x = 0.75
    for n, t, d in [("1", "L'enfant pose sa carte", "sur le lecteur du comptoir"),
                    ("2", "Le logiciel répond", "en moins d'une seconde, avec un son"),
                    ("3", "Vous lisez la couleur", "vert = il entre · rouge = il n'entre pas")]:
        box(s, x, 2.3, 3.75, 2.5, fill=WHITE, outline=BORDER)
        pastille(s, x + 0.35, 2.65, 0.7, GREEN_MAIN, n, size=19)
        text(s, x + 1.25, 2.8, 2.25, 0.45, t, size=17, bold=True)
        text(s, x + 0.35, 3.65, 3.05, 0.8, d, size=14.5, color=MUTED, line=1.2)
        x += 4.0
    box(s, 0.75, 5.2, 11.8, 1.65, fill=GREEN_DARK)
    text(s, 1.2, 5.5, 10.9, 0.5, "Vous n'avez rien à taper. Rien à cliquer.",
         size=25, bold=True, color=YELLOW)
    text(s, 1.2, 6.12, 10.9, 0.45,
         "Si le lecteur ne répond pas, tapez le numéro de la carte à la main "
         "dans le champ, puis appuyez sur Entrée.", size=14, color=WHITE)
    notes(s, "Faire pratiquer le geste a chacun.")

    # ============================================================== 6. écran vert
    s = slide()
    entete(s, "L'écran vert : l'enfant peut entrer",
           "Quatre informations à lire avant de le laisser passer")
    box(s, 0.75, 2.2, 5.6, 4.4, fill=GREEN_PALE, outline=GREEN_MAIN, lw=2.25)
    text(s, 1.05, 2.6, 5.0, 0.55, "ENTRÉE AUTORISÉE", size=26, bold=True,
         color=GREEN_DARK, align=PP_ALIGN.CENTER)
    text(s, 1.05, 3.35, 5.0, 0.45, "Aminata Diallo", size=21, bold=True,
         align=PP_ALIGN.CENTER)
    text(s, 1.05, 3.88, 5.0, 0.35, "M-0001 · Mensuel 4 entrées", size=14,
         color=MUTED, align=PP_ALIGN.CENTER)
    bx = 1.15
    for v, l in [("3", "ENTRÉES\nRESTANTES"), ("23/09", "EXPIRE LE"),
                 ("7", "AVANT VISITE\nOFFERTE")]:
        box(s, bx, 4.5, 1.55, 1.6, fill=WHITE)
        text(s, bx + 0.1, 4.72, 1.35, 0.5, v, size=23, bold=True,
             color=GREEN_DARK, align=PP_ALIGN.CENTER)
        text(s, bx + 0.1, 5.35, 1.35, 0.65, l, size=9.5, bold=True, color=MUTED,
             align=PP_ALIGN.CENTER, space=0, line=1.1)
        bx += 1.7
    y = 2.3
    for t, d in [("Le nom", "Vérifiez que c'est bien l'enfant devant vous."),
                 ("L'abonnement", "Le forfait dont il dispose."),
                 ("Les entrées restantes", "S'il n'en reste qu'une, prévenez le parent."),
                 ("La date d'expiration", "Si la date approche, proposez le renouvellement.")]:
        pastille(s, 6.8, y, 0.38, GREEN_MAIN, "•", size=17)
        text(s, 7.4, y - 0.02, 5.15, 0.36, t, size=16.5, bold=True)
        text(s, 7.4, y + 0.4, 5.15, 0.5, d, size=13.5, color=MUTED, line=1.2)
        y += 1.05
    box(s, 6.8, 6.0, 5.75, 0.72, fill=YELLOW_PALE)
    text(s, 7.05, 6.22, 5.3, 0.4,
         "« Déjà passé 2 fois aujourd'hui » : vérifiez que ce n'est pas une erreur",
         size=12.5, bold=True)
    notes(s, "Le compteur 'avant visite offerte' n'apparait que si la fidelite est activee.")

    # ============================================================== 7. écran rouge
    s = slide()
    entete(s, "L'écran rouge : l'enfant ne peut pas entrer",
           "Le motif est toujours écrit — lisez-le au parent, il comprendra")
    text(s, 0.9, 2.2, 3.9, 0.3, "CE QUI S'AFFICHE", size=12, bold=True, color=MUTED)
    text(s, 5.0, 2.2, 3.3, 0.3, "CE QUE ÇA VEUT DIRE", size=12, bold=True, color=MUTED)
    text(s, 8.6, 2.2, 3.9, 0.3, "CE QUE VOUS FAITES", size=12, bold=True, color=MUTED)
    y = 2.68
    lignes = [
        ("Carte inconnue", "Elle n'est enregistrée sur aucun enfant",
         "Créez la fiche et donnez-lui la carte"),
        ("Carte bloquée", "La carte a été déclarée perdue",
         "Le parent doit utiliser la nouvelle"),
        ("Abonnement expiré le ...", "La date de fin est dépassée",
         "Vendez un nouvel abonnement"),
        ("Plus d'entrées disponibles", "Le forfait est terminé",
         "Vendez un nouvel abonnement"),
        ("Aucun abonnement actif", "L'enfant n'a rien en cours",
         "Vendez un abonnement depuis sa fiche"),
    ]
    for i, (msg, sens, act) in enumerate(lignes):
        box(s, 0.75, y, 11.8, 0.74, fill=RED_PALE if i % 2 == 0 else WHITE)
        text(s, 0.9, y + 0.22, 3.95, 0.4, msg, size=13.5, bold=True, color=RED)
        text(s, 5.0, y + 0.22, 3.45, 0.4, sens, size=13)
        text(s, 8.6, y + 0.22, 3.85, 0.4, act, size=13, bold=True, color=GREEN_DARK)
        y += 0.82
    text(s, 0.75, y + 0.28, 11.8, 0.45,
         "Ne laissez jamais entrer un enfant refusé sans prévenir le responsable.",
         size=15.5, bold=True)
    notes(s, "Les trois derniers cas se resolvent par une vente : voir les slides suivantes.")

    # ============================================================== 8. fidélité
    s = slide(PURPLE_PALE)
    text(s, 0.75, 0.8, 11.8, 0.85, "La visite offerte", size=40, bold=True, color=PURPLE)
    text(s, 0.75, 1.75, 11.8, 0.5,
         "Le logiciel récompense les enfants fidèles tout seul", size=16)
    box(s, 0.75, 2.75, 5.9, 3.9, fill=WHITE)
    text(s, 1.15, 3.25, 5.1, 0.6, "VISITE OFFERTE !", size=30, bold=True,
         color=PURPLE, align=PP_ALIGN.CENTER)
    text(s, 1.15, 4.15, 5.1, 2.2,
         "Toutes les 10 visites payantes, la 11e visite est gratuite.\n\n"
         "Le logiciel compte tout seul et n'entame pas le forfait de l'enfant.",
         size=16, align=PP_ALIGN.CENTER, line=1.3)
    box(s, 7.05, 2.75, 5.5, 3.9, fill=PURPLE)
    text(s, 7.45, 3.2, 4.7, 0.45, "VOTRE RÔLE", size=16, bold=True, color=YELLOW)
    text(s, 7.45, 3.9, 4.7, 2.4,
         "Annoncez-le au parent avec le sourire :\n\n"
         "« Aujourd'hui c'est offert, c'est la 11e visite d'Aminata ! »\n\n"
         "C'est un moment qui fidélise le client.",
         size=15, color=WHITE, line=1.3)
    notes(s, "Le seuil est reglable par le responsable.")

    # ============================================================== 9. parcours
    s = slide(GREEN_DARK)
    text(s, 0.75, 0.7, 11.8, 0.9, "Inscrire un nouvel abonné", size=40,
         bold=True, color=YELLOW)
    text(s, 0.75, 1.7, 11.8, 0.5,
         "Quatre étapes, deux minutes, sans avoir besoin du gérant",
         size=17, color=GREEN_PALE)
    etapes = [("1", "Créer sa fiche"), ("2", "Lui donner une carte"),
              ("3", "Encaisser"), ("4", "Il entre")]
    x = 0.75
    for n, t in etapes:
        box(s, x, 2.7, 2.75, 2.4, fill=WHITE)
        pastille(s, x + 1.02, 3.0, 0.72, YELLOW, n, lc=TEXT, size=20)
        text(s, x + 0.2, 4.0, 2.35, 0.6, t, size=17, bold=True,
             color=GREEN_DARK, align=PP_ALIGN.CENTER, line=1.15)
        if n != "4":
            fleche = box(s, x + 2.83, 3.72, 0.34, 0.3, fill=YELLOW,
                         shape=MSO_SHAPE.RIGHT_ARROW)
        x += 3.1
    box(s, 0.75, 5.5, 11.8, 1.15, fill=YELLOW)
    text(s, 1.15, 5.78, 11.0, 0.6,
         "Les quatre écrans suivants montrent exactement ce que vous verrez.",
         size=19, bold=True, align=PP_ALIGN.CENTER)
    notes(s, "C'est le coeur du guide : l'agent doit pouvoir le faire seul.")

    # ============================================================== 10. étape 1
    s = slide()
    entete(s, "Étape 1 — Créer la fiche",
           "Menu « Membres » puis le bouton jaune « + Nouveau membre »")
    box(s, 0.75, 2.05, 7.0, 4.5, fill=WHITE, outline=BORDER, radius=0.05, lw=1.5)
    text(s, 1.05, 2.3, 4.0, 0.35, "Nouveau membre", size=17, bold=True, color=GREEN_DARK)
    champ(s, 1.05, 2.85, 3.1, "NOM DE L'ENFANT *", "Fatoumata Barry")
    champ(s, 4.45, 2.85, 3.0, "DATE DE NAISSANCE", "14/03/2020")
    champ(s, 1.05, 3.75, 3.1, "NOM DU PARENT", "Alpha Barry")
    champ(s, 4.45, 3.75, 3.0, "TÉLÉPHONE", "628 44 55 66")
    champ(s, 1.05, 4.65, 6.4, "NOTES", "")
    bouton(s, 5.05, 5.75, 1.1, 0.4, "Annuler", fill=GREEN_PALE, lc=GREEN_DARK)
    bouton(s, 6.3, 5.75, 1.15, 0.4, "Enregistrer")
    y = 2.25
    for n, t, d in [
        ("1", "Le nom de l'enfant suffit", "C'est le seul champ obligatoire."),
        ("2", "La date de naissance", "C'est elle qui fera apparaître l'enfant dans "
         "les anniversaires à proposer."),
        ("3", "Le téléphone du parent", "Sans lui, impossible d'envoyer les relances. "
         "Ne le sautez jamais."),
        ("4", "Enregistrer", "La fiche s'ouvre aussitôt : vous enchaînez sur la carte."),
    ]:
        etape(s, 8.15, y, n, t, d, couleur=YELLOW, lc=TEXT, wt=3.9)
        y += 1.12
    notes(s, "Le code membre (M-0005) est attribue automatiquement.")

    # ============================================================== 11. étape 2
    s = slide()
    entete(s, "Étape 2 — Lui donner une carte",
           "Sur la fiche qui vient de s'ouvrir : « + Attribuer une carte »")
    box(s, 0.75, 2.05, 7.0, 4.0, fill=WHITE, outline=BORDER, radius=0.05, lw=1.5)
    text(s, 1.05, 2.35, 5.0, 0.35, "Attribuer une carte RFID", size=17,
         bold=True, color=GREEN_DARK)
    text(s, 1.05, 2.8, 6.0, 0.3,
         "Passez la carte sur le lecteur — le numéro se remplit tout seul.",
         size=12, color=MUTED)
    text(s, 1.05, 3.3, 3.0, 0.2, "NUMÉRO DE LA CARTE *", size=9.5, bold=True, color=MUTED)
    box(s, 1.05, 3.55, 6.4, 0.72, fill=GREEN_PALE, outline=GREEN_MAIN,
        dashed=True, radius=0.15, lw=2.25)
    text(s, 1.05, 3.78, 6.4, 0.3, "En attente de la carte…", size=15, bold=True,
         color=GREEN_DARK, align=PP_ALIGN.CENTER)
    champ(s, 1.05, 4.5, 6.4, "NUMÉRO IMPRIMÉ SUR LA CARTE (FACULTATIF)", "042")
    bouton(s, 5.05, 5.45, 1.1, 0.4, "Annuler", fill=GREEN_PALE, lc=GREEN_DARK)
    bouton(s, 6.3, 5.45, 1.15, 0.4, "Attribuer")
    y = 2.25
    for n, t, d in [
        ("1", "Posez la carte neuve sur le lecteur",
         "Le champ vert se remplit tout seul. Ne tapez rien au clavier."),
        ("2", "Le numéro imprimé est facultatif",
         "Utile si vos cartes portent un numéro visible."),
        ("3", "Cliquez sur « Attribuer »",
         "La carte est active immédiatement."),
    ]:
        etape(s, 8.15, y, n, t, d, wt=3.9)
        y += 1.35
    box(s, 8.15, 6.05, 4.4, 0.75, fill=GREEN_PALE)
    text(s, 8.4, 6.25, 3.9, 0.4, "Un enfant = une seule carte active",
         size=13, bold=True, color=GREEN_DARK)
    notes(s, "Si le logiciel refuse : l'enfant a deja une carte, utiliser 'Carte perdue'.")

    # ============================================================== 12. étape 3
    s = slide()
    entete(s, "Étape 3 — Encaisser",
           "Sur la fiche : bouton jaune « Vendre un abonnement »")
    box(s, 0.75, 2.05, 7.0, 4.4, fill=WHITE, outline=BORDER, radius=0.05, lw=1.5)
    text(s, 1.05, 2.3, 5.5, 0.35, "Vendre un abonnement", size=17, bold=True,
         color=GREEN_DARK)
    fx = 1.05
    for nom, prix, sel in [("Mensuel 4", "180 000", False),
                           ("Mensuel 8", "320 000", True),
                           ("VIP Illimité", "500 000", False)]:
        box(s, fx, 2.8, 2.05, 1.0,
            fill=GREEN_PALE if sel else WHITE,
            outline=GREEN_MAIN if sel else BORDER, lw=2.25 if sel else 1.25)
        text(s, fx + 0.1, 3.0, 1.85, 0.25, nom, size=11.5, bold=True,
             color=GREEN_DARK, align=PP_ALIGN.CENTER, space=0)
        text(s, fx + 0.1, 3.35, 1.85, 0.25, prix + " GNF", size=12, bold=True,
             align=PP_ALIGN.CENTER, space=0)
        fx += 2.2
    champ(s, 1.05, 4.05, 3.1, "MODE DE PAIEMENT", "Orange Money")
    champ(s, 4.45, 4.05, 3.0, "MONTANT PAYÉ (GNF)", "320 000", vc=MUTED)
    box(s, 4.45, 4.95, 3.0, 0.42, fill=YELLOW_PALE, radius=0.2)
    text(s, 4.6, 5.06, 2.75, 0.25, "Prix du catalogue — non modifiable",
         size=9.5, bold=True, space=0)
    bouton(s, 4.15, 5.65, 3.3, 0.45, "Encaisser et imprimer le reçu",
           fill=GREEN_DARK, size=12)
    y = 2.25
    for n, t, d in [
        ("1", "Choisissez la formule", "Elle se met en vert quand elle est sélectionnée."),
        ("2", "Choisissez le paiement", "Espèces, Orange Money, MTN MoMo ou carte."),
        ("3", "Le montant est bloqué", "Vous ne pouvez pas le changer : c'est le prix "
         "du parc. Cela vous protège."),
        ("4", "Encaissez", "Le reçu s'imprime tout seul, avec votre nom dessus."),
    ]:
        etape(s, 8.15, y, n, t, d, couleur=YELLOW, lc=TEXT, wt=3.9)
        y += 1.12
    notes(s, "Rappeler : le montant bloque protege l'employe autant que le parc.")

    # ============================================================== 13. étape 4
    s = slide()
    entete(s, "Étape 4 — L'enfant entre",
           "Il pose sa carte tout de suite : c'est déjà actif")
    box(s, 0.75, 2.2, 6.2, 4.4, fill=GREEN_PALE, outline=GREEN_MAIN, lw=2.25)
    text(s, 1.05, 2.85, 5.6, 0.6, "ENTRÉE AUTORISÉE", size=29, bold=True,
         color=GREEN_DARK, align=PP_ALIGN.CENTER)
    text(s, 1.05, 3.75, 5.6, 0.5, "Fatoumata Barry", size=24, bold=True,
         align=PP_ALIGN.CENTER)
    text(s, 1.05, 4.38, 5.6, 0.35, "M-0005 · Mensuel 8 entrées", size=15,
         color=MUTED, align=PP_ALIGN.CENTER)
    box(s, 2.45, 4.95, 2.8, 1.3, fill=WHITE)
    text(s, 2.55, 5.15, 2.6, 0.55, "7", size=30, bold=True, color=GREEN_DARK,
         align=PP_ALIGN.CENTER)
    text(s, 2.55, 5.8, 2.6, 0.28, "ENTRÉES RESTANTES", size=10.5, bold=True,
         color=MUTED, align=PP_ALIGN.CENTER, space=0)
    text(s, 7.35, 2.4, 5.2, 0.55, "C'est terminé.", size=32, bold=True, color=GREEN_DARK)
    text(s, 7.35, 3.35, 5.2, 3.2,
         "Le parent est inscrit, il a payé, son enfant joue.\n\n"
         "Vous avez fait les quatre étapes seul, en deux minutes, "
         "sans appeler personne.\n\n"
         "L'abonnement est décompté automatiquement à chaque passage : "
         "il vous reste juste à poser la carte, les prochaines fois.",
         size=16.5, line=1.4)
    notes(s, "Fin du parcours d'inscription. Enchainer sur la carte perdue.")

    # ============================================================== 14. carte perdue
    s = slide()
    entete(s, "Une carte perdue ? On la remplace",
           "Ouvrez la fiche de l'enfant : bouton rouge « Carte perdue — la remplacer »")
    box(s, 0.75, 2.05, 11.8, 1.0, fill=RED_PALE, outline=RED, lw=2.25)
    text(s, 1.15, 2.28, 11.0, 0.5,
         "Ne créez JAMAIS une deuxième fiche pour le même enfant.",
         size=20, bold=True, color=RED)
    text(s, 1.15, 2.7, 11.0, 0.3,
         "Vous perdriez son abonnement, ses entrées restantes et son historique.",
         size=13.5)
    box(s, 0.75, 3.35, 7.0, 3.15, fill=WHITE, outline=BORDER, radius=0.05, lw=1.5)
    text(s, 1.05, 3.6, 5.5, 0.35, "Carte perdue — la remplacer", size=16,
         bold=True, color=RED)
    text(s, 1.05, 4.02, 6.2, 0.3,
         "L'ancienne carte sera désactivée automatiquement.", size=11.5, color=MUTED)
    box(s, 1.05, 4.45, 6.4, 0.62, fill=GREEN_PALE, outline=GREEN_MAIN,
        dashed=True, radius=0.15, lw=2.25)
    text(s, 1.05, 4.62, 6.4, 0.3, "Posez la NOUVELLE carte…", size=14, bold=True,
         color=GREEN_DARK, align=PP_ALIGN.CENTER)
    champ(s, 1.05, 5.25, 3.1, "MOTIF", "Carte perdue")
    bouton(s, 5.35, 5.65, 2.1, 0.42, "Remplacer la carte", fill=RED)
    y = 3.5
    for n, t, d in [
        ("1", "Ouvrez la fiche de l'enfant", "Menu Membres, puis « Ouvrir »."),
        ("2", "Posez la nouvelle carte", "Le champ vert se remplit tout seul."),
        ("3", "Validez", "L'ancienne carte est refusée dès cet instant. "
         "L'abonnement et les entrées restantes sont conservés."),
    ]:
        etape(s, 8.15, y, n, t, d, couleur=RED, wt=3.9)
        y += 1.05
    notes(s, "Si le parent retrouve l'ancienne carte : la reposer sur la fiche, "
             "elle se reactive.")

    # ============================================================== 15. relances
    s = slide()
    entete(s, "Appeler les parents",
           "Menu « Relances WhatsApp » — le chiffre indique combien de parents appeler")
    box(s, 0.75, 2.15, 11.8, 1.75, fill=GREEN_PALE)
    text(s, 1.1, 2.4, 5.0, 0.35, "Abonnements à renouveler", size=15, bold=True,
         color=GREEN_DARK)
    text(s, 1.1, 2.95, 3.0, 0.3, "Mariama Toure", size=13, bold=True)
    text(s, 4.3, 2.95, 3.0, 0.3, "Expire le 02/09/2026", size=12.5, color=MUTED)
    bouton(s, 9.6, 2.9, 1.6, 0.42, "WhatsApp", fill=GREEN_MAIN)
    text(s, 1.1, 3.4, 10.0, 0.3,
         "Un clic ouvre WhatsApp avec le message déjà écrit et le bon numéro.",
         size=12, color=MUTED)
    box(s, 0.75, 4.1, 11.8, 1.75, fill=PURPLE_PALE)
    text(s, 1.1, 4.35, 5.0, 0.35, "Anniversaires — 30 prochains jours", size=15,
         bold=True, color=PURPLE)
    text(s, 1.1, 4.9, 3.0, 0.3, "Aminata Diallo", size=13, bold=True)
    text(s, 4.3, 4.9, 3.0, 0.3, "7 ans dans 6 jours", size=12.5, color=MUTED)
    bouton(s, 9.6, 4.85, 1.6, 0.42, "WhatsApp", fill=GREEN_MAIN)
    text(s, 1.1, 5.35, 10.0, 0.3,
         "Proposez une fête d'anniversaire au parc — c'est une vente facile.",
         size=12, color=MUTED)
    box(s, 0.75, 6.05, 11.8, 0.8, fill=YELLOW)
    text(s, 1.15, 6.3, 11.0, 0.4,
         "Relisez toujours le message avant d'envoyer, et ajoutez un mot personnel.",
         size=15, bold=True, align=PP_ALIGN.CENTER)
    notes(s, "Faire cette tache a un moment calme de la journee.")

    # ============================================================== 16. règles
    s = slide(GREEN_DARK)
    text(s, 0.75, 0.65, 11.8, 0.85, "Les 5 règles d'or", size=40, bold=True, color=YELLOW)
    y = 1.95
    for i, (t, d) in enumerate([
        ("Chaque action porte votre nom", "entrées, encaissements : tout est enregistré"),
        ("Ne prêtez jamais votre compte", "même à un collègue, même cinq minutes"),
        ("Les prix ne se négocient pas", "le montant vient du catalogue, il est bloqué"),
        ("Une carte = un enfant", "carte perdue : on remplace, on ne recrée pas la fiche"),
        ("Ne fermez pas la fenêtre noire", "c'est le moteur du logiciel"),
    ], 1):
        pastille(s, 0.85, y, 0.7, YELLOW, str(i), lc=TEXT, size=19)
        text(s, 1.85, y + 0.03, 10.4, 0.4, t, size=20, bold=True, color=WHITE)
        text(s, 1.85, y + 0.5, 10.4, 0.35, d, size=14.5, color=GREEN_PALE)
        y += 1.08
    notes(s, "A afficher au comptoir.")

    # ============================================================== 17. problèmes
    s = slide()
    entete(s, "En cas de problème", "Trois pannes courantes, trois solutions")
    y = 2.25
    for t, d in [
        ("L'écran ne s'ouvre pas",
         "Vérifiez que la fenêtre « Didikids Parc - Serveur » est ouverte dans la "
         "barre des tâches. Sinon, double-cliquez sur l'icône du Bureau."),
        ("Le lecteur ne réagit pas",
         "Débranchez puis rebranchez le lecteur USB et relancez « Lecteur RFID ». "
         "En attendant, tapez le numéro de la carte à la main."),
        ("Un menu a disparu",
         "Ce n'est pas une panne : chacun voit uniquement les pages qui le "
         "concernent. Le reste est réservé au responsable."),
    ]:
        box(s, 0.75, y, 11.8, 1.25, fill=GREEN_PALE)
        text(s, 1.1, y + 0.24, 3.6, 0.45, t, size=17, bold=True, color=GREEN_DARK)
        text(s, 4.9, y + 0.22, 7.4, 0.9, d, size=14, line=1.2)
        y += 1.45
    box(s, 0.75, 6.35, 11.8, 0.85, fill=YELLOW)
    text(s, 1.1, 6.63, 11.1, 0.45,
         "En cas de doute, appelez le responsable. Aucune question n'est bête.",
         size=17, bold=True, align=PP_ALIGN.CENTER)
    notes(s, "Laisser le numero du responsable affiche au comptoir.")

    # ============================================================== 18. fin
    s = slide(GREEN_DARK)
    text(s, 0.75, 2.4, 11.8, 0.95, "Des questions ?", size=46, bold=True,
         color=WHITE, align=PP_ALIGN.CENTER)
    box(s, 3.4, 3.95, 6.5, 1.55, fill=YELLOW)
    text(s, 3.7, 4.3, 5.9, 0.85, "Posez la carte.\nLisez la couleur.",
         size=24, bold=True, align=PP_ALIGN.CENTER, line=1.2)
    text(s, 0.75, 5.95, 11.8, 0.5,
         "Didikids Parc — parce que chaque enfant mérite de s'épanouir",
         size=16, color=GREEN_PALE, align=PP_ALIGN.CENTER)
    notes(s, "Terminer par une mise en pratique sur le vrai lecteur.")

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "Didikids_Parc_Guide_Agent.pptx")
    prs.save(out)
    return out


if __name__ == "__main__":
    print("OK :", build())
